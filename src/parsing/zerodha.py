"""Zerodha tradebook parser.

Normalizes Zerodha's Equity tradebook XLSX (NSE/BSE) into the canonical
schema. See docs/design/parsing_architecture.md §6 for the observed format.

Zerodha tradebooks contain only BUY/SELL rows in INR with no FX rate and no
total amount — both are derived here. Dividends arrive in a separate report
and are handled by a future ``ZerodhaDividendParser`` (not yet implemented).
"""

from typing import Optional

import openpyxl
import pandas as pd

from src.parsing.base import StatementParser, ParsedStatement, ParseError, CANONICAL_COLUMNS
from src.ticker_utils import fetch_inr_per_eur


# Header cells that identify a Zerodha tradebook. The header row sits a few rows
# below the top (blank rows + "Client ID" + report title precede it).
_TRADEBOOK_MARKERS = {"Symbol", "ISIN", "Trade Date", "Exchange", "Trade Type", "Quantity", "Price"}

# Zerodha -> canonical column mapping (column indices are irrelevant; we match
# on the header cell text).
_COLUMN_MAP = {
    "Symbol": "Ticker",
    "ISIN": "ISIN",
    "Trade Date": "Date",
    "Trade Type": "Type",
    "Quantity": "Quantity",
    "Price": "Price per share",
}


def _find_header_row(rows):
    """Return the index of the tradebook header row, or None if not found."""
    for idx, row in enumerate(rows):
        cells = {str(c).strip() for c in row if c is not None}
        if _TRADEBOOK_MARKERS.issubset(cells):
            return idx
    return None


def _read_rows(path: str):
    """Read all rows from the Equity sheet of ``path``.

    Returns (header_row, data_rows) or raises ParseError.
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface any read failure
        raise ParseError(f"Could not open Zerodha workbook: {exc}") from exc

    try:
        sheet_name = "Equity" if "Equity" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    header_idx = _find_header_row(rows)
    if header_idx is None:
        raise ParseError(
            "Could not find Zerodha tradebook header (expected Symbol/ISIN/Trade Date/Trade Type)"
        )
    return rows[header_idx], rows[header_idx + 1 :]


class ZerodhaParser(StatementParser):
    source = "zerodha"

    def can_parse(self, path: str) -> bool:
        if not path.lower().endswith(".xlsx"):
            return False
        try:
            header, _ = _read_rows(path)
        except Exception:  # noqa: BLE001 - a non-Zerodha file simply won't match
            return False
        cells = {str(c).strip() for c in header if c is not None}
        return _TRADEBOOK_MARKERS.issubset(cells)

    def parse(self, path: str, inr_per_eur: Optional[float] = None) -> ParsedStatement:
        header, data_rows = _read_rows(path)

        # INR->EUR rate: use an explicit override if given, otherwise fetch the
        # current rate from yfinance (EURINR=X).
        if inr_per_eur is None or inr_per_eur <= 0:
            inr_per_eur = fetch_inr_per_eur()
        if inr_per_eur is None or inr_per_eur <= 0:
            raise ParseError(
                "Could not determine an INR→EUR rate (EURINR=X fetch failed); "
                "provide inr_per_eur explicitly."
            )

        # Build a header -> column-index lookup from the header row.
        col_index = {}
        for i, cell in enumerate(header):
            if cell is not None:
                key = str(cell).strip()
                if key in _COLUMN_MAP:
                    col_index[_COLUMN_MAP[key]] = i

        required = {"Ticker", "Date", "Type", "Quantity", "Price per share"}
        missing = required - set(col_index)
        if missing:
            raise ParseError(f"Zerodha tradebook missing columns: {sorted(missing)}")

        records = []
        warnings = []
        c = col_index
        for row_num, row in enumerate(data_rows, start=1):
            # Data rows have a blank leading column; the Symbol sits at index 1.
            symbol = row[c["Ticker"]] if len(row) > c["Ticker"] else None
            if symbol is None or str(symbol).strip() == "":
                continue

            trade_type = str(row[c["Type"]]).strip().lower() if row[c["Type"]] is not None else ""
            if trade_type not in ("buy", "sell"):
                warnings.append(f"Skipped row {row_num}: unsupported trade type '{trade_type}'")
                continue

            try:
                quantity = float(row[c["Quantity"]])
                price = float(row[c["Price per share"]])
            except (TypeError, ValueError):
                warnings.append(f"Skipped row {row_num}: unparseable quantity/price")
                continue

            records.append(
                {
                    "Ticker": str(symbol).strip().upper(),
                    "ISIN": row[c["ISIN"]] if "ISIN" in c else None,
                    "Date": row[c["Date"]],
                    "Type": "BUY" if trade_type == "buy" else "SELL",
                    "Quantity": quantity,
                    "Price per share": price,
                    "Total Amount": quantity * price,
                    "Currency": "INR",
                    "FX Rate": inr_per_eur,
                    "Source": self.source,
                }
            )

        if not records:
            raise ParseError("Zerodha tradebook contained no trade rows")

        df = pd.DataFrame(records)

        # Normalize dates robustly (handles datetime, ISO strings, and NaT).
        df["Date"] = df["Date"].apply(
            lambda x: pd.to_datetime(str(x).replace("Z", ""), utc=True, errors="coerce")
        )

        # Ensure canonical column order (pad any optional columns we didn't set).
        for col in CANONICAL_COLUMNS:
            if col not in df.columns:
                df[col] = None
        df = df[CANONICAL_COLUMNS]

        return ParsedStatement(
            source=self.source,
            rows=len(df),
            data=df,
            warnings=warnings,
        )